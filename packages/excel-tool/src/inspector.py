#!/usr/bin/env python3
"""
Excel 智能检查器
- 自动识别表头行数与多层表头（处理合并单元格与总标题行）
- 提取复合表头结构
- 敏感列特征与数据正则智能嗅探（身份证、手机号、姓名、邮箱、银行卡等）
- 提取前 15 行样例数据
"""

import sys
import os
import json
import re
import csv
from typing import Dict, List, Any, Optional

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

# ── 敏感列识别规则 ─────────────────────────────────────────────────────────────

SENSITIVE_PATTERNS = [
    {
        "type": "id_card",
        "name_re": re.compile(r"身份证|证件号|公民身份|id_?card|identity|id_?num", re.I),
        "data_re": re.compile(r"^[1-9]\d{5}(?:18|19|20)\d{2}(?:0[1-9]|1[0-2])(?:0[1-9]|[12]\d|3[01])\d{3}[\dXx]$"),
        "rule": "id_card_mask",
        "label": "18位身份证号",
        "desc": "保留前6位行政区划与后4位，中间脱敏"
    },
    {
        "type": "phone",
        "name_re": re.compile(r"手机|电话|联系电话|联系方式|phone|mobile|tel", re.I),
        "data_re": re.compile(r"^1[3-9]\d{9}$"),
        "rule": "phone_mask",
        "label": "手机号码",
        "desc": "保留前3位与后4位，中间4位掩码"
    },
    {
        "type": "email",
        "name_re": re.compile(r"邮箱|邮件|电子信箱|email|mail", re.I),
        "data_re": re.compile(r"^[\w\.-]+@[\w\.-]+\.[a-zA-Z]{2,}$"),
        "rule": "email_mask",
        "label": "电子邮箱",
        "desc": "保留首字母与域名"
    },
    {
        "type": "bank_card",
        "name_re": re.compile(r"银行卡|卡号|银行账号|bank_?card|account_?no", re.I),
        "data_re": re.compile(r"^[1-9]\d{15,18}$"),
        "rule": "bank_card_mask",
        "label": "银行卡号",
        "desc": "保留前4与后4位"
    },
    {
        "type": "name",
        "name_re": re.compile(r"^(?:员工|客户|用户|患者|学生|旅客|联系人)?(?:姓名|名字|name|contacts?)$", re.I),
        "data_re": re.compile(r"^[\u4e00-\u9fa5]{2,4}$"),
        "rule": "name_mask",
        "label": "姓名",
        "desc": "保留姓氏，名字部分掩码"
    },
    {
        "type": "amount",
        "name_re": re.compile(r"工资|薪资|薪酬|底薪|实发|基本工资|奖金|收入|报酬|salary|wage|compensation", re.I),
        "data_re": None,
        "rule": "amount_mask",
        "label": "敏感薪资",
        "desc": "数值掩码/区间化"
    }
]


def resolve_merged_matrix(ws: Worksheet, max_scan_row: int) -> List[List[Any]]:
    """展开合并单元格，构建展开后的前 max_scan_row 行单元格二维矩阵"""
    matrix = []
    max_col = ws.max_column or 1
    for r in range(1, max_scan_row + 1):
        row_vals = []
        for c in range(1, max_col + 1):
            val = ws.cell(row=r, column=c).value
            row_vals.append(str(val).strip() if val is not None else "")
        matrix.append(row_vals)

    # 处理合并单元格填充
    for rng in ws.merged_cells.ranges:
        min_col, min_row, max_c, max_r = rng.bounds
        if min_row > max_scan_row:
            continue
        top_val = ws.cell(row=min_row, column=min_col).value
        fill_val = str(top_val).strip() if top_val is not None else ""
        end_r = min(max_r, max_scan_row)
        for r in range(min_row, end_r + 1):
            for c in range(min_col, max_c + 1):
                if r <= len(matrix) and c <= len(matrix[r - 1]):
                    matrix[r - 1][c - 1] = fill_val

    return matrix


def detect_headers(matrix: List[List[str]]) -> Dict[str, Any]:
    """
    智能推断表头起始行、结束行与复合列名
    """
    total_rows = len(matrix)
    if total_rows == 0:
        return {"header_start": 1, "header_end": 1, "headers": [], "levels": 1}

    max_cols = max(len(r) for r in matrix) if matrix else 0

    # 1. 查找是否第一行是跨越多列的大标题（例如跨列合并，或者第一行所有非空单元格都相同）
    header_start_idx = 0
    first_row_vals = [c for c in matrix[0] if c]
    if len(set(first_row_vals)) == 1 and max_cols > 1 and total_rows > 1:
        # 第一行整行由同一个合并单元格或者单一标题填充
        header_start_idx = 1

    # 2. 判断表头层数（最多检测到 3 层）
    # 观察 header_start_idx 之后的行，如果后一行包含许多重复值或者下一行依然像表头
    header_end_idx = header_start_idx
    for check_idx in range(header_start_idx + 1, min(header_start_idx + 3, total_rows)):
        current_row = matrix[check_idx]
        prev_row = matrix[check_idx - 1]
        # 如果当前行大多是文本且不全是数字/日期，且上一行有合并/重复值，可能是多层表头
        has_duplicate_in_prev = len(set(c for c in prev_row if c)) < len([c for c in prev_row if c])
        non_empty_current = len([c for c in current_row if c])
        if has_duplicate_in_prev and non_empty_current > 0:
            header_end_idx = check_idx
        else:
            break

    # 3. 构造复合列名
    composite_headers = []
    header_levels = (header_end_idx - header_start_idx) + 1

    for c_idx in range(max_cols):
        parts = []
        for r_idx in range(header_start_idx, header_end_idx + 1):
            val = matrix[r_idx][c_idx] if c_idx < len(matrix[r_idx]) else ""
            if val and val not in parts:
                parts.append(val)
        col_name = " / ".join(parts) if parts else f"Column_{c_idx + 1}"
        composite_headers.append(col_name)

    return {
        "header_start": header_start_idx + 1,  # 1-indexed
        "header_end": header_end_idx + 1,      # 1-indexed
        "headers": composite_headers,
        "levels": header_levels,
    }


def sniff_sensitive_column(col_name: str, samples: List[Any]) -> Optional[Dict[str, Any]]:
    """基于列名和样例数据特征识别敏感类型"""
    # 负向关键词：如果是部门、岗位、状态等，不应识别为人名
    neg_name_re = re.compile(r"部门|组织|公司|机构|岗位|职位|职务|级别|状态|类型|类别|行业|城市|省份|地址|国家|科目|备注|描述", re.I)

    # 1. 优先通过列名匹配
    for p in SENSITIVE_PATTERNS:
        if p["name_re"].search(col_name):
            if p["type"] == "name" and neg_name_re.search(col_name):
                continue
            return {
                "column": col_name,
                "type": p["type"],
                "rule": p["rule"],
                "label": p["label"],
                "reason": f"列名包含「{p['label']}」关键词",
                "desc": p["desc"],
            }

    # 2. 列名未匹配时，通过非空样本数据正则检查
    non_empty_samples = [str(s).strip() for s in samples if s is not None and str(s).strip()]
    if non_empty_samples:
        for p in SENSITIVE_PATTERNS:
            if not p["data_re"]:
                continue
            # 如果是人名检测，但列名有负向关键词，跳过
            if p["type"] == "name" and neg_name_re.search(col_name):
                continue
            matched_count = sum(1 for s in non_empty_samples if p["data_re"].match(s))
            if matched_count >= max(1, len(non_empty_samples) * 0.5):
                return {
                    "column": col_name,
                    "type": p["type"],
                    "rule": p["rule"],
                    "label": p["label"],
                    "reason": f"数据样本高度符合「{p['label']}」格式",
                    "desc": p["desc"],
                }

    return None


def detect_csv_encoding(filepath: str) -> str:
    """自动嗅探 CSV 编码（优先尝试 utf-8-sig、utf-8、gb18030、gbk 等）"""
    for enc in ['utf-8-sig', 'utf-8', 'gb18030', 'gbk', 'big5', 'latin1']:
        try:
            with open(filepath, 'r', encoding=enc) as f:
                f.read(8192)
            return enc
        except Exception:
            continue
    return 'utf-8'


def inspect_csv(filepath: str, max_sample_rows: int = 15) -> Dict[str, Any]:
    """检查 CSV 文件结构、推断表头并嗅探敏感字段"""
    enc = detect_csv_encoding(filepath)
    all_rows: List[List[str]] = []
    with open(filepath, 'r', encoding=enc, errors='replace') as f:
        reader = csv.reader(f)
        for row in reader:
            all_rows.append([str(c).strip() for c in row])

    if not all_rows:
        raise ValueError("CSV 文件为空")

    scan_rows = min(10, len(all_rows))
    matrix = all_rows[:scan_rows]
    header_info = detect_headers(matrix)

    headers = header_info["headers"]
    header_end = header_info["header_end"]
    total_rows = len(all_rows)
    total_cols = len(headers)

    data_start_row = header_end + 1
    sample_rows = []
    col_samples: Dict[str, List[Any]] = {h: [] for h in headers}

    curr_row_idx = header_end
    while curr_row_idx < total_rows and len(sample_rows) < max_sample_rows:
        row_vals = all_rows[curr_row_idx]
        row_dict = {}
        has_content = False
        for c_idx, h in enumerate(headers):
            val_str = row_vals[c_idx] if c_idx < len(row_vals) else ""
            if val_str:
                has_content = True
            row_dict[h] = val_str
            col_samples[h].append(val_str)

        if has_content:
            sample_rows.append(row_dict)
        curr_row_idx += 1

    sensitive_columns = []
    for h in headers:
        detection = sniff_sensitive_column(h, col_samples.get(h, []))
        if detection:
            sensitive_columns.append(detection)

    fname = os.path.basename(filepath)
    sheet_name = os.path.splitext(fname)[0] or "CSV"

    return {
        "filename": fname,
        "filepath": os.path.abspath(filepath),
        "fileSizeBytes": os.path.getsize(filepath),
        "sheets": [sheet_name],
        "activeSheet": sheet_name,
        "rowCount": max(0, total_rows - header_end),
        "columnCount": total_cols,
        "headerLevels": header_info["levels"],
        "headerStartRow": header_info["header_start"],
        "headerEndRow": header_info["header_end"],
        "dataStartRow": data_start_row,
        "headers": headers,
        "sampleRows": sample_rows,
        "sensitiveColumns": sensitive_columns,
    }


def inspect_excel(filepath: str, max_sample_rows: int = 15) -> Dict[str, Any]:
    if not os.path.exists(filepath):
        raise FileNotFoundError(f"文件不存在: {filepath}")

    if filepath.lower().endswith('.csv'):
        return inspect_csv(filepath, max_sample_rows)

    wb = openpyxl.load_workbook(filepath, data_only=True)
    sheet_names = wb.sheetnames
    if not sheet_names:
        raise ValueError("Excel 文件中无有效工作表")

    ws = wb.active or wb[sheet_names[0]]

    # 扫描前 10 行用于推断表头
    scan_rows = min(10, ws.max_row or 1)
    matrix = resolve_merged_matrix(ws, scan_rows)
    header_info = detect_headers(matrix)

    headers = header_info["headers"]
    header_end = header_info["header_end"]
    total_rows = ws.max_row or 0
    total_cols = len(headers)

    # 提取样例数据 (从 header_end + 1 开始)
    data_start_row = header_end + 1
    sample_rows = []
    col_samples: Dict[str, List[Any]] = {h: [] for h in headers}

    curr_row = data_start_row
    while curr_row <= total_rows and len(sample_rows) < max_sample_rows:
        row_dict = {}
        has_content = False
        for c_idx, h in enumerate(headers, start=1):
            val = ws.cell(row=curr_row, column=c_idx).value
            if val is not None:
                has_content = True
                val_str = str(val).strip()
            else:
                val_str = ""
            row_dict[h] = val_str
            col_samples[h].append(val_str)

        if has_content:
            sample_rows.append(row_dict)
        curr_row += 1

    # 敏感列识别
    sensitive_columns = []
    for h in headers:
        detection = sniff_sensitive_column(h, col_samples.get(h, []))
        if detection:
            sensitive_columns.append(detection)

    return {
        "filename": os.path.basename(filepath),
        "filepath": os.path.abspath(filepath),
        "fileSizeBytes": os.path.getsize(filepath),
        "sheets": sheet_names,
        "activeSheet": ws.title,
        "rowCount": max(0, total_rows - header_end),
        "columnCount": total_cols,
        "headerLevels": header_info["levels"],
        "headerStartRow": header_info["header_start"],
        "headerEndRow": header_info["header_end"],
        "dataStartRow": data_start_row,
        "headers": headers,
        "sampleRows": sample_rows,
        "sensitiveColumns": sensitive_columns,
    }


def main():
    # 强制标准输出与标准错误使用 UTF-8 编码，防止 Windows 平台默认 GBK 代码页乱码
    if hasattr(sys.stdout, 'reconfigure'):
        try:
            sys.stdout.reconfigure(encoding='utf-8')
        except Exception:
            pass
    if hasattr(sys.stderr, 'reconfigure'):
        try:
            sys.stderr.reconfigure(encoding='utf-8')
        except Exception:
            pass

    if len(sys.argv) < 2:
        print(json.dumps({"error": "缺少文件路径参数"}))
        sys.exit(1)

    filepath = sys.argv[1]
    try:
        result = inspect_excel(filepath)
        print(json.dumps({"success": True, "data": result}, ensure_ascii=False))
    except Exception as e:
        import traceback
        print(json.dumps({
            "success": False,
            "error": str(e),
            "traceback": traceback.format_exc()
        }, ensure_ascii=False))
        sys.exit(1)


if __name__ == "__main__":
    main()
